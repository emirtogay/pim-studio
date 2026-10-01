"""Reading somebody's list of product codes.

Two things are worth pinning here and the rest is plumbing.

**The column is found by evidence, not by its header.** These files arrive with
headers in Turkish, or missing, or lying, and with three plausible-looking code
columns of which one is a barcode. Scoring each column by how many of its
values match products we hold is what survives all of that.

**Nothing is invented.** A token that matches no option comes back in
`not_found` in the spelling it was written in. The alternative — parsing
`2001524000` into a style and a colour by guessing where the boundary is —
produces a ModelColor that looks entirely correct and belongs to no product.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from pim_studio.reports.uploads import (
    parse_csv,
    parse_pasted,
    parse_upload,
    parse_workbook,
)

KNOWN = ["L3558.132", "TH2568.144", "2011169.-", "60KFA0791.83J", "AH4463.132"]


def sheet_bytes(rows: list[list[object]], title: str = "Sheet1") -> bytes:
    book = Workbook()
    book.active.title = title
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# ── the five spellings ─────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "written",
    [
        "L3558.132",  # Omnitron, dotted
        "L3558_132",  # BrandLibrary, underscored
        "L3558 132",  # retyped
        "L3558132",  # order file, concatenated
        "l3558.132",  # lower case
        "  L3558.132  ",  # pasted out of a cell
        "L3558.132.T4",  # size-suffixed article code
        "L3558.132.TXL",  # letter size
    ],
)
def test_every_spelling_resolves_to_the_same_option(written):
    assert parse_pasted(written, KNOWN).model_colors == ["L3558.132"]


def test_a_code_that_matches_nothing_is_returned_as_written():
    """Not normalised, not corrected — the user has to recognise it to act
    on it."""
    parsed = parse_pasted("zz_9999", KNOWN)
    assert parsed.model_colors == []
    assert parsed.not_found == ["zz_9999"]


def test_a_dash_size_code_is_not_truncated():
    """Watches and accessories carry `-` as their size. `2011169.-` is already
    a ModelColor and stripping a suffix from it invents a different one."""
    assert parse_pasted("2011169.-", KNOWN).model_colors == ["2011169.-"]


def test_duplicates_collapse_across_spellings():
    parsed = parse_pasted("L3558.132\nL3558_132\nL3558.132.T4", KNOWN)
    assert parsed.model_colors == ["L3558.132"]


def test_input_order_is_preserved():
    """An uploaded list is usually already in the order somebody wants to read
    it; re-sorting makes them find their place again."""
    parsed = parse_pasted("AH4463.132, L3558.132, TH2568.144", KNOWN)
    assert parsed.model_colors == ["AH4463.132", "L3558.132", "TH2568.144"]


@pytest.mark.parametrize("separator", ["\n", ",", ";", "\t", "\r\n"])
def test_pasted_text_splits_on_every_separator_people_use(separator):
    text = separator.join(["L3558.132", "TH2568.144"])
    assert len(parse_pasted(text, KNOWN).model_colors) == 2


def test_a_space_does_not_split_a_code():
    """`L3558 132` is one code written with a space, not two codes."""
    assert parse_pasted("L3558 132", KNOWN).model_colors == ["L3558.132"]


# ── finding the column ─────────────────────────────────────────────────────


def test_the_column_with_the_most_matches_wins_over_a_promising_header():
    """`sku` sounds right and holds nothing we know; `Poz.` sounds like a line
    number and holds the codes. The scores decide."""
    data = sheet_bytes(
        [
            ["sku", "Poz.", "Qty"],
            ["99999", "L3558.132", 4],
            ["99998", "TH2568.144", 2],
            ["99997", "AH4463.132", 9],
        ]
    )
    parsed = parse_workbook(data, KNOWN)
    assert parsed.column == "Poz."
    assert len(parsed.model_colors) == 3


def test_a_numeric_column_never_wins():
    data = sheet_bytes(
        [["Qty", "Code"], [1312132, "L3558.132"], [7318144, "TH2568.144"]]
    )
    assert parse_workbook(data, KNOWN).column == "Code"


def test_the_list_is_found_on_whichever_sheet_holds_it():
    """It is rarely the first sheet — it is behind a cover sheet, or on the one
    named after the season."""
    book = Workbook()
    book.active.title = "Cover"
    book.active.append(["26SS season opening"])
    codes = book.create_sheet("Order")
    codes.append(["Article"])
    for code in ("L3558.132", "TH2568.144"):
        codes.append([code])
    buffer = io.BytesIO()
    book.save(buffer)

    parsed = parse_workbook(buffer.getvalue(), KNOWN)
    assert parsed.sheet == "Order"
    assert len(parsed.model_colors) == 2


def test_a_file_holding_no_codes_we_know_explains_itself_rather_than_guessing():
    data = sheet_bytes([["Name", "City"], ["Alice", "Warsaw"], ["Bob", "Prague"]])
    parsed = parse_workbook(data, KNOWN)
    assert parsed.model_colors == []
    assert "No column" in parsed.problem


# ── csv ────────────────────────────────────────────────────────────────────


def test_semicolons_are_read_as_a_delimiter():
    """Which is what a Polish Excel writes."""
    parsed = parse_csv(b"nr;kod;ilosc\n1;L3558.132;4\n2;TH2568.144;2\n", KNOWN)
    assert len(parsed.model_colors) == 2


def test_a_byte_order_mark_does_not_break_the_first_header():
    parsed = parse_csv("﻿code\nL3558.132\n".encode(), KNOWN)
    assert parsed.model_colors == ["L3558.132"]


def test_an_empty_file_says_so():
    assert "empty" in parse_csv(b"", KNOWN).problem


# ── dispatch ───────────────────────────────────────────────────────────────


def test_an_unreadable_extension_is_refused_with_the_alternative():
    with pytest.raises(ValueError, match="paste the codes"):
        parse_upload("scan.pdf", b"%PDF", KNOWN)


@pytest.mark.parametrize("name", ["list.csv", "list.tsv", "list.txt"])
def test_text_extensions_route_to_the_csv_reader(name):
    parsed = parse_upload(name, b"code\nL3558.132\n", KNOWN)
    assert parsed.model_colors == ["L3558.132"]


def test_xlsx_routes_to_the_workbook_reader():
    data = sheet_bytes([["Article"], ["L3558.132"]])
    assert parse_upload("list.xlsx", data, KNOWN).model_colors == ["L3558.132"]
