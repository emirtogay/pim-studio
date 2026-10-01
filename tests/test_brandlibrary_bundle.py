"""The handoff bundle — the module's actual output.

Our Omnitron connection is read-only, so a photo on the photo server is not a
photo on a product page until somebody in Turkey uploads it. These tests are
about the thing they receive.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from pim_studio.brandlibrary import bundle as bundle_module
from pim_studio.models import Base, HandoffBundle, PublishedImage


@pytest.fixture
def db(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'bundle.db'}", future=True)
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine, expire_on_commit=False)()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def publish(db, model_color: str, asset: str, edition: str = "original") -> PublishedImage:
    suffix = "" if edition == "original" else "_ZAL"
    row = PublishedImage(
        model_color=model_color,
        brandlibrary_key=model_color.replace(".", "_"),
        asset_name=asset,
        edition=edition,
        remote_name=f"{asset}{suffix}.jpg",
        public_url=f"https://photos.example/pim-studio/{asset}{suffix}.jpg",
        bytes=1024,
    )
    db.add(row)
    db.commit()
    return row


def build(db, tmp_path, **kwargs):
    return bundle_module.build(db, tmp_path / "bundles", fetch=False, **kwargs)


# ── what goes in ───────────────────────────────────────────────────────────


def test_nothing_published_produces_no_bundle(db, tmp_path) -> None:
    """A normal weekly outcome. An empty ZIP is a thing somebody has to open
    to discover it is empty."""
    assert build(db, tmp_path) is None
    assert db.scalars(select(HandoffBundle)).all() == []


def test_a_bundle_carries_every_published_photo(db, tmp_path) -> None:
    publish(db, "AF0276.J2G", "AF0276_J2G_32")
    publish(db, "AF0276.J2G", "AF0276_J2G_31")
    publish(db, "L3558.132", "L3558_132_01")

    result = build(db, tmp_path)

    assert result is not None
    assert result.files == 3
    assert result.options == 2


def test_a_photo_joins_exactly_one_bundle(db, tmp_path) -> None:
    """Re-sending is not harmless — it reads to the recipient as a correction."""
    publish(db, "AF0276.J2G", "AF0276_J2G_32")

    first = build(db, tmp_path)
    second = build(db, tmp_path)

    assert first is not None
    assert second is None


def test_a_photo_published_after_a_bundle_lands_in_the_next_one(db, tmp_path) -> None:
    publish(db, "AF0276.J2G", "AF0276_J2G_32")
    build(db, tmp_path)
    publish(db, "L3558.132", "L3558_132_01")

    second = build(db, tmp_path)

    assert second is not None
    assert second.files == 1


# ── the Excel, which is the operative half ─────────────────────────────────


def test_the_sheet_lists_every_photo_with_its_public_url(db, tmp_path) -> None:
    """Omnitron imports images by URL. This column is what does the work."""
    publish(db, "AF0276.J2G", "AF0276_J2G_32")
    publish(db, "L3558.132", "L3558_132_01")

    result = build(db, tmp_path)

    sheet = load_workbook(result.xlsx).active
    header = [c.value for c in sheet[6]]
    assert header == ["ModelColor", "Edition", "File name", "Image URL", "Published"]

    urls = {sheet.cell(row=r, column=4).value for r in (7, 8)}
    assert urls == {
        "https://photos.example/pim-studio/AF0276_J2G_32.jpg",
        "https://photos.example/pim-studio/L3558_132_01.jpg",
    }


def test_the_two_editions_are_separate_rows(db, tmp_path) -> None:
    publish(db, "AF0276.J2G", "AF0276_J2G_32", edition="original")
    publish(db, "AF0276.J2G", "AF0276_J2G_32", edition="zalando")

    result = build(db, tmp_path)

    sheet = load_workbook(result.xlsx).active
    editions = sorted(sheet.cell(row=r, column=2).value for r in (7, 8))
    assert editions == ["original", "zalando"]
    names = sorted(sheet.cell(row=r, column=3).value for r in (7, 8))
    assert names == ["AF0276_J2G_32.jpg", "AF0276_J2G_32_ZAL.jpg"]


def test_the_sheet_says_what_it_is_for(db, tmp_path) -> None:
    """It is emailed and read by people who were not in the conversation."""
    publish(db, "AF0276.J2G", "AF0276_J2G_32")

    sheet = load_workbook(build(db, tmp_path).xlsx).active

    assert "no image in Omnitron" in str(sheet.cell(row=2, column=1).value)
    assert "1 options" in str(sheet.cell(row=4, column=1).value)


def test_the_header_stays_visible_while_scrolling(db, tmp_path) -> None:
    """A thousand rows of URLs with no column names is unreadable."""
    publish(db, "AF0276.J2G", "AF0276_J2G_32")

    sheet = load_workbook(build(db, tmp_path).xlsx).active

    assert sheet.freeze_panes == "A7"
    assert sheet.auto_filter.ref is not None


# ── the ZIP ────────────────────────────────────────────────────────────────


def test_the_zip_names_every_file_the_sheet_names(db, tmp_path) -> None:
    publish(db, "AF0276.J2G", "AF0276_J2G_32")
    publish(db, "AF0276.J2G", "AF0276_J2G_32", edition="zalando")

    result = build(db, tmp_path)

    with zipfile.ZipFile(result.zip) as archive:
        assert sorted(archive.namelist()) == [
            "AF0276_J2G_32.jpg",
            "AF0276_J2G_32_ZAL.jpg",
        ]


def test_a_bundle_can_be_built_without_a_zip(db, tmp_path) -> None:
    publish(db, "AF0276.J2G", "AF0276_J2G_32")

    result = build(db, tmp_path, with_zip=False)

    assert result.zip is None
    assert result.xlsx.exists()


# ── the row outlives the files ─────────────────────────────────────────────


def test_pruning_clears_files_and_keeps_the_record(db, tmp_path) -> None:
    """"What did we hand over in August" outlives the ZIP by years."""
    for n in range(3):
        publish(db, f"A{n}.1", f"A{n}_1_01")
        build(db, tmp_path)

    cleared = bundle_module.prune(db, tmp_path / "bundles", keep=1)

    assert cleared == 2
    bundles = db.scalars(select(HandoffBundle).order_by(HandoffBundle.id)).all()
    assert [b.status for b in bundles] == ["cleared", "cleared", "ready"]
    # The counts are still true after the files are gone.
    assert all(b.file_count == 1 for b in bundles)
    assert not Path(bundles[0].xlsx_path).exists()
    assert Path(bundles[-1].xlsx_path).exists()


def test_pruning_nothing_when_there_is_little_to_prune(db, tmp_path) -> None:
    publish(db, "AF0276.J2G", "AF0276_J2G_32")
    build(db, tmp_path)

    assert bundle_module.prune(db, tmp_path / "bundles", keep=10) == 0


def test_sending_is_recorded(db, tmp_path) -> None:
    """Null means built and not yet sent, which is a real state."""
    publish(db, "AF0276.J2G", "AF0276_J2G_32")
    result = build(db, tmp_path)
    assert db.get(HandoffBundle, result.bundle_id).sent_at is None

    bundle_module.mark_sent(db, result.bundle_id)

    assert db.get(HandoffBundle, result.bundle_id).sent_at is not None
