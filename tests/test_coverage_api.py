"""The coverage endpoints.

The property worth pinning hardest is that the downloaded workbook and the
on-screen numbers come from the same bytes. Both paths read `coverage.json`,
so a test that asserts the endpoints agree with `build_report` over the same
extract is the test that keeps a spreadsheet from quietly disagreeing with the
page it was downloaded from.

The rest is about the upload: an unmatched code has to survive the round trip,
because an upload that reports on 380 of 500 codes and says nothing is worse
than one that fails.
"""

from __future__ import annotations

import io
import json
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from pim_studio.admin import app as app_module
from pim_studio.admin import coverage_api
from pim_studio.admin.auth import current_user
from pim_studio.db import get_session, session_scope
from pim_studio.models import Attribute, AttributeValue, Product, User
from pim_studio.reports.coverage import EXTRACT_VERSION
from pim_studio.reports.values import TooManyValues

pytestmark = pytest.mark.usefixtures("temp_db")


def _extract() -> dict:
    """A three-option catalogue in the shape `to_extract` emits.

    Written out by hand rather than generated from a database, so a change to
    the extract format has to be made here too — which is the point. The
    frontend reads this same shape and cannot be caught by a Python test.
    """
    return {
        "version": EXTRACT_VERSION,
        "generatedAt": datetime(2026, 8, 21, 6, 0, tzinfo=UTC).isoformat(),
        "markets": ["PL", "CZ", "SK", "HU"],
        "shared": "*",
        "attributes": [
            {
                "key": "product_name",
                "label": "Product name",
                "owner": "us",
                "group": "Content",
                "perMarket": True,
                "appliesWhen": "",
                "sourced": "CZ,PL",
            },
            {
                "key": "price",
                "label": "Price",
                "owner": "partner",
                "group": "Commercial",
                "perMarket": True,
                "appliesWhen": "",
                "sourced": "CZ,PL",
            },
            {
                "key": "image_0",
                "label": "First image",
                "owner": "system",
                "group": "Images",
                "perMarket": False,
                "appliesWhen": "",
                "sourced": "*",
            },
            {
                "key": "sleeve_type",
                "label": "Sleeve type",
                "owner": "partner",
                "group": "Classification",
                "perMarket": False,
                "appliesWhen": "Polo,Shirt",
                "sourced": "*",
            },
        ],
        "issues": [],
        "profiles": [],
        "stockCaveat": "provisional",
        "unread": [["product_name", "SK"], ["product_name", "HU"]],
        "products": [
            {
                "mc": "A.1",
                "name": "Complete polo",
                "season": "26SS",
                "sub": "Polo",
                "gender": "MAN",
                "skus": 5,
                "markets": "CZ,PL",
                "flags": [],
                "filled": {
                    "product_name": "CZ,PL",
                    "price": "CZ,PL",
                    "image_0": "*",
                    "sleeve_type": "*",
                },
                "comment": "",
                "image": "https://example.test/a.jpg",
                "stock": 12,
            },
            {
                "mc": "B.2",
                "name": "Empty polo",
                "season": "26SS",
                "sub": "Polo",
                "gender": "MAN",
                "skus": 3,
                "markets": "PL",
                "flags": ["no_image"],
                "filled": {"price": "PL"},
                "comment": "",
                "image": "",
                "stock": 0,
            },
            {
                "mc": "C.3",
                "name": "A wallet",
                "season": "OUTLET",
                "sub": "Wallet",
                "gender": "UNISEX",
                "skus": 1,
                "markets": "",
                "flags": ["removed", "not_in_catalogue"],
                "filled": {"product_name": "PL", "image_0": "*"},
                "comment": "REMOVED",
                "image": "https://example.test/c.jpg",
                "stock": 0,
            },
        ],
    }


@pytest.fixture
def client(tmp_path, monkeypatch) -> TestClient:
    demo = tmp_path / "demo"
    demo.mkdir()
    (demo / "coverage.json").write_text(json.dumps(_extract()), encoding="utf-8")
    monkeypatch.setattr(coverage_api, "_demo_dir", lambda: demo)

    app_module.app.dependency_overrides[current_user] = lambda: User(
        id=1, email="t@example.com", role="planning", name="T", is_active=True
    )
    try:
        yield TestClient(app_module.app, raise_server_exceptions=False)
    finally:
        app_module.app.dependency_overrides.clear()


@pytest.fixture
def no_extract(tmp_path, monkeypatch) -> TestClient:
    monkeypatch.setattr(coverage_api, "_demo_dir", lambda: tmp_path / "nothing")
    app_module.app.dependency_overrides[current_user] = lambda: User(
        id=1, email="t@example.com", role="planning", name="T", is_active=True
    )
    try:
        yield TestClient(app_module.app, raise_server_exceptions=False)
    finally:
        app_module.app.dependency_overrides.clear()


PROFILE = {
    "label": "Test",
    "requirements": [
        {"attribute": "product_name", "markets": ["PL", "CZ"]},
        {"attribute": "image_0", "markets": []},
        {"attribute": "sleeve_type", "markets": []},
    ],
}


# ── the extract ────────────────────────────────────────────────────────────


def test_extract_is_served(client):
    response = client.get("/api/coverage")
    assert response.status_code == 200
    assert len(response.json()["products"]) == 3


def test_missing_extract_says_how_to_make_one(no_extract):
    response = no_extract.get("/api/coverage")
    assert response.status_code == 503
    assert "build_demo_data_from_db" in response.json()["detail"]


def test_presets_are_listed_without_downloading_the_extract(client):
    keys = [p["key"] for p in client.get("/api/coverage/profiles").json()["profiles"]]
    assert "sellable" in keys and "localised" in keys


def test_endpoints_require_a_session(tmp_path, monkeypatch):
    monkeypatch.setattr(coverage_api, "_demo_dir", lambda: tmp_path)
    app_module.app.dependency_overrides.clear()
    anonymous = TestClient(app_module.app, raise_server_exceptions=False)
    assert anonymous.get("/api/coverage").status_code == 401
    assert anonymous.post("/api/coverage/summary", json=PROFILE).status_code == 401


# ── summary ────────────────────────────────────────────────────────────────


def test_summary_counts_the_three_standings(client):
    body = client.post("/api/coverage/summary", json=PROFILE).json()
    assert body["scope"]["population"] == 3
    assert body["standing"]["complete"] == 1  # A.1
    assert body["standing"]["incomplete"] == 2  # B.2 and C.3


def test_summary_keeps_unknown_out_of_the_fill_rate(client):
    """`product_name` was never read for SK or HU, and the profile does not ask
    for them — but the requirement rows must still never invent a 0%."""
    body = client.post("/api/coverage/summary", json=PROFILE).json()
    rows = {(r["attribute"], r["market"]): r for r in body["requirements"]}
    assert rows[("product_name", "CZ")]["missing"] == 2
    assert rows[("product_name", "CZ")]["fillRate"] == pytest.approx(1 / 3)


def test_summary_marks_an_inapplicable_attribute_na_not_missing(client):
    body = client.post("/api/coverage/summary", json=PROFILE).json()
    sleeve = next(r for r in body["requirements"] if r["attribute"] == "sleeve_type")
    assert sleeve["na"] == 1  # the wallet
    assert sleeve["missing"] == 1  # B.2, a polo with nothing filled


def test_summary_routes_work_to_the_owner_who_may_write_it(client):
    work = client.post("/api/coverage/summary", json=PROFILE).json()["work"]
    assert work["partner"]["gaps"] == 1
    assert work["us"]["gaps"] == 3
    assert work["system"]["gaps"] == 1


def test_a_preset_key_is_accepted_instead_of_an_explicit_list(client):
    body = client.post("/api/coverage/summary", json={"profile": "localised"}).json()
    assert body["profile"]["key"] == "localised"


def test_an_empty_definition_is_refused_rather_than_answered(client):
    """It would report every option complete — a plausible number and a
    meaningless one."""
    response = client.post("/api/coverage/summary", json={"requirements": []})
    assert response.status_code == 400
    assert "at least one attribute" in response.json()["detail"]


def test_a_stale_extract_version_is_refused(client, tmp_path, monkeypatch):
    demo = tmp_path / "old"
    demo.mkdir()
    stale = {**_extract(), "version": EXTRACT_VERSION + 1}
    (demo / "coverage.json").write_text(json.dumps(stale), encoding="utf-8")
    monkeypatch.setattr(coverage_api, "_demo_dir", lambda: demo)
    response = client.post("/api/coverage/summary", json=PROFILE)
    assert response.status_code == 503
    assert "Regenerate" in response.json()["detail"]


# ── scoping ────────────────────────────────────────────────────────────────


def test_scope_narrows_to_an_uploaded_list(client):
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "modelColors": ["A.1"]}
    ).json()
    assert body["scope"]["population"] == 1
    assert body["standing"]["complete"] == 1


def test_unmatched_codes_come_back_rather_than_vanishing(client):
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "modelColors": ["A.1", "NOPE.1"]}
    ).json()
    assert body["scope"]["notFound"] == ["NOPE.1"]
    assert body["scope"]["population"] == 1


def test_scope_can_narrow_to_a_market_catalogue(client):
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "inCatalogue": "CZ"}
    ).json()
    assert body["scope"]["population"] == 1
    assert "CZ catalogue" in body["scope"]["description"]


def test_scope_can_narrow_by_season(client):
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "season": ["OUTLET"]}
    ).json()
    assert body["scope"]["population"] == 1



def test_scope_can_narrow_to_options_carrying_a_price(client):
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "priced": "any"}
    ).json()
    assert body["scope"]["population"] == 2
    assert "priced somewhere" in body["scope"]["description"]


def test_price_scope_is_per_market(client):
    """B.2 has a Polish price and no Czech one, so it belongs in a CZ price
    list and not in a report of what CZ already covers."""
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "priced": "CZ"}
    ).json()
    assert body["scope"]["population"] == 1
    assert "priced in CZ" in body["scope"]["description"]


def test_stock_scope_takes_no_market(client):
    """One warehouse behind four markets. Refused rather than quietly widened
    to `any`, which would answer a different question than the one asked."""
    response = client.post(
        "/api/coverage/summary", json={**PROFILE, "stocked": "PL"}
    )
    assert response.status_code == 422
    assert "one pool of stock" in response.text


def test_scope_can_narrow_to_options_in_stock(client):
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "stocked": "any"}
    ).json()
    assert body["scope"]["population"] == 1


def test_an_unknown_price_market_is_refused_not_ignored(client):
    response = client.post("/api/coverage/summary", json={**PROFILE, "priced": "DE"})
    assert response.status_code == 422


def test_issue_flags_reach_the_server_now(client):
    """The filter the download used to drop on the floor: the page showed one
    population and the workbook was built from a wider one."""
    body = client.post(
        "/api/coverage/summary", json={**PROFILE, "flags": ["no_image"]}
    ).json()
    assert body["scope"]["population"] == 1
    assert "No image" in body["scope"]["description"]


def test_an_unknown_issue_key_is_refused(client):
    response = client.post(
        "/api/coverage/summary", json={**PROFILE, "flags": ["no_unicorn"]}
    )
    assert response.status_code == 422


def test_ticking_article_mcs_is_a_column_and_does_not_change_the_grain():
    """The regression. `base_code_format` is a real attribute — `seed.py`
    registers it as "Article (MCS)" at sku grain — and it was briefly used as
    the sentinel for the size switch. Ticking that row in the grid therefore
    flipped the whole sheet to one row per SKU *and* swallowed the column the
    reader had asked for, neither of which they requested."""
    from pim_studio.admin.coverage_api import ScopeBody

    body = ScopeBody(
        values=[
            {"attribute": "base_code_format", "markets": []},
            {"attribute": "product_name", "markets": ["PL"]},
        ]
    )
    assert body.by_size() is False
    assert [s.attribute for s in body.selections()] == [
        "base_code_format",
        "product_name",
    ]


def test_the_grain_is_its_own_field():
    from pim_studio.admin.coverage_api import ScopeBody

    body = ScopeBody(
        bySize=True, values=[{"attribute": "product_name", "markets": ["PL"]}]
    )
    assert body.by_size() is True
    assert [s.attribute for s in body.selections()] == ["product_name"]


def test_the_grain_defaults_to_options():
    from pim_studio.admin.coverage_api import ScopeBody

    assert ScopeBody().by_size() is False

# ── resolving an upload ────────────────────────────────────────────────────


def test_pasted_codes_resolve_across_spellings(client):
    body = client.post(
        "/api/coverage/resolve/text", json={"text": "a.1\nB_2\nC.3.T04"}
    ).json()
    assert body["modelColors"] == ["A.1", "B.2", "C.3"]
    assert body["notFound"] == []


def test_pasted_codes_report_what_did_not_match(client):
    body = client.post("/api/coverage/resolve/text", json={"text": "A.1, ZZ.9"}).json()
    assert body["modelColors"] == ["A.1"]
    assert body["notFound"] == ["ZZ.9"]


def test_a_list_matching_nothing_explains_itself(client):
    body = client.post("/api/coverage/resolve/text", json={"text": "ZZ.9\nYY.8"}).json()
    assert body["matched"] == 0
    assert "not synced" in body["problem"]


def _workbook_bytes(rows: list[list[str]]) -> bytes:
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_an_uploaded_workbook_finds_the_column_holding_codes(client):
    data = _workbook_bytes(
        [
            ["Line", "Article", "Qty"],
            ["1", "A.1.T4", "10"],
            ["2", "B_2", "4"],
            ["3", "ZZ.9", "1"],
        ]
    )
    response = client.post(
        "/api/coverage/resolve",
        files={"file": ("order.xlsx", data, "application/vnd.ms-excel")},
    )
    body = response.json()
    assert body["column"] == "Article"
    assert body["modelColors"] == ["A.1", "B.2"]
    assert body["notFound"] == ["ZZ.9"]


def test_an_uploaded_csv_is_read(client):
    response = client.post(
        "/api/coverage/resolve",
        files={"file": ("list.csv", b"code\nA.1\nC.3\n", "text/csv")},
    )
    assert response.json()["modelColors"] == ["A.1", "C.3"]


def test_an_unreadable_file_type_is_refused_with_advice(client):
    response = client.post(
        "/api/coverage/resolve",
        files={"file": ("photo.png", b"\x89PNG", "image/png")},
    )
    assert response.status_code == 400
    assert "paste the codes" in response.json()["detail"]


def test_a_corrupt_workbook_is_the_users_problem_not_a_500(client):
    response = client.post(
        "/api/coverage/resolve",
        files={"file": ("broken.xlsx", b"not a workbook", "application/vnd.ms-excel")},
    )
    assert response.status_code == 400


def test_resolve_needs_something_to_resolve(client):
    response = client.post("/api/coverage/resolve", data={})
    assert response.status_code == 400


# ── the workbook download ──────────────────────────────────────────────────


def test_report_downloads_a_named_workbook(client):
    response = client.post("/api/coverage/report", json={"profile": "sellable"})
    assert response.status_code == 200
    assert "Northwind_coverage_sellable_" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert "Summary" in book.sheetnames


def test_report_and_summary_agree_on_the_population(client):
    """The one property that makes a download trustworthy: page and file read
    the same bytes, so they cannot drift."""
    scope = {**PROFILE, "modelColors": ["A.1", "B.2", "NOPE.1"]}
    summary = client.post("/api/coverage/summary", json=scope).json()
    download = client.post("/api/coverage/report", json=scope)
    assert download.headers["x-coverage-options"] == str(summary["scope"]["population"])
    assert download.headers["x-coverage-not-found"] == str(
        len(summary["scope"]["notFound"])
    )


def test_report_carries_the_not_found_sheet(client):
    response = client.post(
        "/api/coverage/report", json={**PROFILE, "modelColors": ["A.1", "NOPE.1"]}
    )
    book = load_workbook(io.BytesIO(response.content))
    assert "Not found" in book.sheetnames


# ── the work queue ─────────────────────────────────────────────────────────


def test_queue_lists_each_gap_against_the_person_who_fixes_it(client):
    queues = client.post("/api/coverage/queue", json=PROFILE).json()["queues"]
    assert {(g["modelColor"], g["scope"]) for g in queues["partner"]} == {("B.2", "ALL")}
    assert {g["modelColor"] for g in queues["us"]} == {"B.2", "C.3"}


def test_queue_respects_the_scope(client):
    queues = client.post(
        "/api/coverage/queue", json={**PROFILE, "modelColors": ["A.1"]}
    ).json()["queues"]
    assert all(not rows for rows in queues.values())


# ── the Values sheet ───────────────────────────────────────────────────────
#
# A database appears in a file that is otherwise entirely about the extract,
# which is worth a sentence. The extract carries presence and never text, so a
# reader who wants to read the descriptions themselves can only be served from
# the mirror. What is pinned here is the boundary that keeps that safe: the
# extract still decides which options the file is about, and the mirror only
# fills cells for options it has already chosen.


def _seed_value(key: str, model_color: str, text: str, market: str = "PL") -> None:
    with session_scope() as db:
        attribute = db.scalar(select(Attribute).where(Attribute.key == key))
        if attribute is None:
            attribute = Attribute(
                key=key,
                label=key.replace("_", " ").title(),
                write_channel="omnitron_import",
                value_grain="model_color",
                per_market=True,
            )
            db.add(attribute)
            db.flush()
        db.add(
            Product(
                sku=f"{model_color}.T1",
                base_code_format=f"{model_color}.T1",
                model_color=model_color,
                base_code=model_color.split(".")[0],
            )
        )
        db.add(
            AttributeValue(
                entity_kind="model_color",
                entity_key=model_color,
                attribute_id=attribute.id,
                market=market,
                value=text,
                state="live",
                state_lane="current",
                entry_grain="model_color",
                source_kind="omnitron",
            )
        )


def _seed_attribute(key: str) -> None:
    """The registry row and nothing else. In production `seed.py` guarantees
    these exist; a throwaway database has to be told."""
    with session_scope() as db:
        db.add(
            Attribute(
                key=key,
                label=key.replace("_", " ").title(),
                write_channel="omnitron_import",
                value_grain="model_color",
                per_market=True,
            )
        )


def _sheets(response) -> list[str]:
    return load_workbook(io.BytesIO(response.content)).sheetnames


def test_a_report_without_values_still_has_four_sheets_and_no_values_tab(client):
    response = client.post("/api/coverage/report", json=PROFILE)
    assert response.status_code == 200
    assert "Values" not in _sheets(response)
    assert "x-coverage-values-read-at" not in response.headers


def test_a_report_without_values_never_touches_the_database(client):
    """Every request the panel sends today ticks no attributes. None of them
    should open a connection, which is what keeps the old path exactly as it
    was."""

    class Explodes:
        def execute(self, *a, **k):  # noqa: ANN002, ANN003
            raise AssertionError("the report opened a database session")

        scalars = execute
        rollback = execute

    app_module.app.dependency_overrides[get_session] = lambda: Explodes()
    try:
        assert client.post("/api/coverage/report", json=PROFILE).status_code == 200
    finally:
        app_module.app.dependency_overrides.pop(get_session, None)


def test_asking_for_values_adds_the_sheet(client):
    _seed_value("product_name", "AAA.1", "Koszulka Polo")
    response = client.post(
        "/api/coverage/report",
        json={**PROFILE, "values": [{"attribute": "product_name", "markets": ["PL"]}]},
    )
    assert response.status_code == 200
    assert "Values" in _sheets(response)
    assert response.headers["x-coverage-values-read-at"]

    sheet = load_workbook(io.BytesIO(response.content))["Values"]
    assert [c.value for c in sheet[4]][:3] == [
        "ModelColor",
        "Name",
        "Product Name · PL",
    ]


def test_the_values_sheet_never_adds_a_row_the_summary_did_not_count(client):
    """The load-bearing test. `Z.9` is in the mirror and not in the extract, so
    a Values sheet built from the database rather than scoped by the extract
    would quietly widen the population of the file."""
    _seed_value("product_name", "Z.9", "Not in the extract")
    response = client.post(
        "/api/coverage/report",
        json={**PROFILE, "values": [{"attribute": "product_name", "markets": ["PL"]}]},
    )
    assert response.status_code == 200

    book = load_workbook(io.BytesIO(response.content))
    codes = {row[0].value for row in book["Values"].iter_rows(min_row=5) if row[0].value}
    assert "Z.9" not in codes
    assert codes == {
        row[0].value
        for row in book["ModelColor Detail"].iter_rows(min_row=5)
        if row[0].value
    }
    assert len(codes) == int(response.headers["x-coverage-options"])


def test_values_against_an_empty_mirror_is_a_sheet_of_blanks_not_a_500(client):
    """The honest failure mode. The attribute is in the registry but nothing
    has ever been synced into it, so the column comes back empty — which is a
    finding, not an error."""
    _seed_attribute("product_name")
    response = client.post(
        "/api/coverage/report",
        json={**PROFILE, "values": [{"attribute": "product_name", "markets": ["PL"]}]},
    )
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content))["Values"]
    assert sheet.max_column == 3
    assert all(row[2].value in (None, "") for row in sheet.iter_rows(min_row=5))


def test_an_unknown_attribute_is_ignored_rather_than_a_400(client):
    response = client.post(
        "/api/coverage/report",
        json={**PROFILE, "values": [{"attribute": "no_such_thing", "markets": []}]},
    )
    assert response.status_code == 200
    assert "Values" not in _sheets(response)


def test_too_many_values_is_a_413_that_says_how_to_narrow_it(client, monkeypatch):
    _seed_value("product_name", "AAA.1", "Koszulka Polo")
    monkeypatch.setattr(coverage_api, "load_values", _refuses)
    response = client.post(
        "/api/coverage/report",
        json={**PROFILE, "values": [{"attribute": "product_name", "markets": ["PL"]}]},
    )
    assert response.status_code == 413
    assert "Narrow the scope" in response.json()["detail"]


def _refuses(*args, **kwargs):  # noqa: ANN002, ANN003
    raise TooManyValues(options=18_036, columns=5, limit=50_000)
