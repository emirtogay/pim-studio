"""Reading a Values sheet back — the round trip has to be exact.

A sheet the Coverage page wrote, edited in one cell, must come back as one
edit. Everything else the parser does is about not inventing edits: a blank is
not a clearing, an unrecognised header is ignored and named, a market comes
only from the header, and sizes that disagree on a per-option value produce no
edit at all.
"""

from __future__ import annotations

import io

from openpyxl import Workbook

from pim_studio.reports.values_upload import CellEdit, parse_table, parse_values_upload

REGISTRY = {
    "product_name": ("Product name", True),
    "description": ("Description", True),
    "pl_sezon": ("Season (PL team)", False),
}


def sheet_bytes(rows: list[list[object]], title: str = "Values") -> bytes:
    """A workbook shaped like the Coverage export: title, subtitle, header."""
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(["Values"])
    sheet.append(["Values read live from the mirror at …"])
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_a_values_sheet_with_one_edited_cell_yields_one_edit() -> None:
    data = sheet_bytes([
        ["ModelColor", "Name", "Product name · PL", "Product name · CZ", "Season (PL team)"],
        ["L3558.132", "Polo", "Koszulka polo", "", "26AW"],
        ["PH9619.166", "Polo", "", "", ""],
    ])
    parsed = parse_values_upload("values.xlsx", data, REGISTRY)
    assert parsed.problem is None
    assert parsed.sheet == "Values" and parsed.header_row == 3 and parsed.rows_read == 2
    assert parsed.edits == [
        CellEdit("L3558.132", "product_name", "PL", "Koszulka polo"),
        CellEdit("L3558.132", "pl_sezon", "*", "26AW"),
    ]
    assert parsed.columns == ["product_name · PL", "product_name · CZ", "pl_sezon · *"]
    assert parsed.unknown_columns == []


def test_headers_may_use_the_key_or_the_label_and_unknown_ones_are_named() -> None:
    table = [
        ["ModelColor", "product_name · pl", "Description (HU)", "Nazwa · PL", "Product name"],
        ["L3558.132", "Koszulka", "<p>Leírás</p>", "x", "y"],
    ]
    parsed = parse_table(table, REGISTRY)
    assert parsed.edits == [
        CellEdit("L3558.132", "product_name", "PL", "Koszulka"),
        CellEdit("L3558.132", "description", "HU", "<p>Leírás</p>"),
    ]
    # `Nazwa · PL` is nobody's label; `Product name` names no market for a
    # per-market attribute. Both are ignored and reported, never guessed.
    assert parsed.unknown_columns == ["Nazwa · PL", "Product name"]


def test_a_size_grain_sheet_folds_agreeing_sizes_and_refuses_disagreeing_ones() -> None:
    table = [
        ["base_code_format", "SKU", "ModelColor", "Size", "Name", "Product name · PL"],
        ["L3558.132.T2", "L3558.132.T2", "L3558.132", "T2", "Polo", "Koszulka"],
        ["L3558.132.T3", "L3558.132.T3", "L3558.132", "T3", "Polo", "Koszulka"],
        ["PH9619.166.T2", "PH9619.166.T2", "PH9619.166", "T2", "Polo", "Jedna"],
        ["PH9619.166.T3", "PH9619.166.T3", "PH9619.166", "T3", "Polo", "Druga"],
    ]
    parsed = parse_table(table, REGISTRY)
    assert parsed.edits == [CellEdit("L3558.132", "product_name", "PL", "Koszulka")]
    assert parsed.sizes_disagree == [("PH9619.166", "product_name", "PL")]


def test_a_sheet_without_a_header_row_is_a_problem_not_a_guess() -> None:
    parsed = parse_table([["a", "b"], ["c", "d"]], REGISTRY)
    assert parsed.edits == [] and parsed.problem and "header row" in parsed.problem


def test_the_named_values_sheet_wins_over_a_cover_sheet() -> None:
    book = Workbook()
    cover = book.active
    cover.title = "Summary"
    cover.append(["Coverage"])
    values = book.create_sheet("Values")
    values.append(["ModelColor", "Product name · SK"])
    values.append(["L3558.132", "Polokošeľa"])
    out = io.BytesIO()
    book.save(out)
    parsed = parse_values_upload("report.xlsx", out.getvalue(), REGISTRY)
    assert parsed.sheet == "Values"
    assert parsed.edits == [CellEdit("L3558.132", "product_name", "SK", "Polokošeľa")]


def test_a_csv_round_trips_too() -> None:
    data = "ModelColor;Product name · HU\nL3558.132;Póló\n".encode()
    parsed = parse_values_upload("values.csv", data, REGISTRY)
    assert parsed.edits == [CellEdit("L3558.132", "product_name", "HU", "Póló")]


def test_a_whole_number_from_excel_is_not_a_float() -> None:
    # `parse_table` takes strings from the CSV path, but the workbook path hands
    # it whatever openpyxl read; `_norm` is what makes 26.0 into "26" either way.
    table: list[list] = [["ModelColor", "Season (PL team)"], ["L3558.132", 26.0]]
    parsed = parse_table(table, REGISTRY)
    assert parsed.edits == [CellEdit("L3558.132", "pl_sezon", "*", "26")]


def test_an_unknown_extension_is_refused_with_advice() -> None:
    parsed = parse_values_upload("values.pdf", b"", REGISTRY)
    assert parsed.problem and ".xlsx or .csv" in parsed.problem
