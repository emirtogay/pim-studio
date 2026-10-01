"""Reading a Values sheet back in — the other half of the round trip.

The Coverage page writes a `Values` sheet (one row per option) or a
`SKU Detail` sheet (one row per size), with one column per attribute and
market headed `Product name · PL`. Somebody edits it in Excel and uploads it
here. This module turns that file into cell edits — one (option, attribute,
market, value) each — and says everything it had to decide along the way:
which sheet it read, where the header row was, which columns it did not
recognise, and which option's sizes disagreed on a value that is stored once
per option.

Three rules keep an edited spreadsheet from becoming a silent bulk write:

**An empty cell is no change.** Clearing a value in Omnitron is a different
write from setting one and this panel does not do it. A blank means "leave
it", which is also what a downloaded sheet with a blank means: we hold
nothing there.

**A header is matched, never guessed.** `Product name · PL` and
`product_name · PL` both resolve; `Nazwa · PL` does not, and is reported as
ignored rather than mapped to the nearest thing. The market comes only from
the header's own suffix.

**Sizes must agree.** On a size-grain sheet a value stored per option appears
once per size. If the sizes of one option say different things, none of them
is taken — the sheet was edited inconsistently and a person has to say which.
"""

from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field

from ..models import MARKET_AGNOSTIC

MAX_ROWS = 50_000
#: How far down a sheet the header row may sit. The Coverage workbook puts a
#: title and a subtitle above it; a hand-made sheet puts it on row 1.
_HEADER_SEARCH_ROWS = 25

#: Columns that identify a row rather than carry a value, in either grain.
_KEY_HEADERS = frozenset(
    {"modelcolor", "sku", "base_code_format", "size", "name", "model color", "mcs"}
)
_MARKETS = ("PL", "CZ", "SK", "HU")

#: `Product name · PL`, `product_name · PL`, `Product name - PL`, `Product name (PL)`.
_SUFFIXED = re.compile(
    r"^(?P<head>.+?)\s*(?:[·|\-–—]\s*|\(\s*)(?P<market>PL|CZ|SK|HU)\s*\)?\s*$", re.IGNORECASE
)


@dataclass(frozen=True)
class CellEdit:
    """One cell somebody changed: an option, an attribute, a market, a value."""

    model_color: str
    attribute: str
    market: str
    value: str


@dataclass
class ParsedValues:
    """What a file turned out to contain, and what was decided reading it."""

    edits: list[CellEdit] = field(default_factory=list)
    sheet: str = ""
    header_row: int = 0
    rows_read: int = 0
    #: Headers that matched no attribute. Reported, never mapped by guesswork.
    unknown_columns: list[str] = field(default_factory=list)
    #: Attribute columns read, as `attribute · market`, so the panel can say
    #: exactly what it will change.
    columns: list[str] = field(default_factory=list)
    #: (model_color, attribute, market) whose sizes disagreed. Dropped.
    sizes_disagree: list[tuple[str, str, str]] = field(default_factory=list)
    problem: str | None = None


@dataclass(frozen=True)
class _Column:
    index: int
    attribute: str
    market: str


def _norm(cell: object) -> str:
    if cell is None:
        return ""
    text = str(cell)
    # Excel round-trips a whole number as `26.0`; a season is not a float.
    if isinstance(cell, float) and cell.is_integer():
        text = str(int(cell))
    return text.strip()


def _find_header(table: list[list[str]]) -> int:
    """The row holding `ModelColor` or `SKU`, or -1."""
    for r, row in enumerate(table[:_HEADER_SEARCH_ROWS]):
        cells = {c.strip().lower() for c in row}
        if "modelcolor" in cells or "model color" in cells or "sku" in cells:
            return r
    return -1


def _resolve_header(
    header: str, by_label: dict[str, tuple[str, bool]], by_key: dict[str, tuple[str, bool]]
) -> tuple[str, str] | None:
    """(attribute key, market) for a value column header, or None."""
    text = header.strip()
    if not text or text.lower() in _KEY_HEADERS:
        return None
    market = MARKET_AGNOSTIC
    head = text
    found = _SUFFIXED.match(text)
    if found:
        head, market = found.group("head").strip(), found.group("market").upper()
    entry = by_key.get(head.lower()) or by_label.get(head.lower())
    if entry is None:
        return None
    key, per_market = entry
    if per_market and market == MARKET_AGNOSTIC:
        # A per-market attribute with no market in the header cannot be placed.
        return None
    if not per_market:
        market = MARKET_AGNOSTIC
    return key, market


def parse_table(
    table: list[list[str]],
    registry: dict[str, tuple[str, bool]],
    *,
    sheet: str = "",
) -> ParsedValues:
    """`registry` maps an attribute key to (label, per_market)."""
    by_key = {key.lower(): (key, per_market) for key, (_label, per_market) in registry.items()}
    by_label = {
        label.lower(): (key, per_market) for key, (label, per_market) in registry.items()
    }
    result = ParsedValues(sheet=sheet)
    header_at = _find_header(table)
    if header_at < 0:
        result.problem = (
            f"{sheet or 'The file'} has no header row naming ModelColor or SKU in "
            f"its first {_HEADER_SEARCH_ROWS} rows."
        )
        return result
    header = [c.strip() for c in table[header_at]]
    result.header_row = header_at + 1
    lowered = [h.lower() for h in header]
    mc_index = next(
        (i for i, h in enumerate(lowered) if h in ("modelcolor", "model color")), -1
    )
    sku_index = next((i for i, h in enumerate(lowered) if h == "sku"), -1)

    columns: list[_Column] = []
    for index, name in enumerate(header):
        if index in (mc_index, sku_index):
            continue
        resolved = _resolve_header(name, by_label, by_key)
        if resolved is None:
            if name and name.lower() not in _KEY_HEADERS:
                result.unknown_columns.append(name)
            continue
        key, market = resolved
        columns.append(_Column(index, key, market))
        result.columns.append(f"{key} · {market}")
    if not columns:
        result.problem = (
            "No column header names an attribute. Headers look like "
            "'Product name · PL' or 'product_name · PL'."
        )
        return result

    # Values per (option, attribute, market), keeping every distinct text seen
    # so a size-grain sheet whose sizes disagree can be refused rather than
    # decided.
    seen: dict[tuple[str, str, str], dict[str, int]] = defaultdict(lambda: defaultdict(int))
    order: list[tuple[str, str, str]] = []
    rows = table[header_at + 1 : header_at + 1 + MAX_ROWS]
    for row in rows:
        model_color = ""
        if mc_index >= 0 and mc_index < len(row):
            model_color = _norm(row[mc_index])
        if not model_color and sku_index >= 0 and sku_index < len(row):
            # A size-grain sheet without a ModelColor column: the option is the
            # SKU with its size suffix removed, which `sync/products.py` derives
            # the same way.
            sku = _norm(row[sku_index])
            model_color = sku.rsplit(".", 1)[0] if sku.count(".") >= 2 else ""
        if not model_color:
            continue
        result.rows_read += 1
        for column in columns:
            if column.index >= len(row):
                continue
            value = _norm(row[column.index])
            if not value:
                continue
            key = (model_color, column.attribute, column.market)
            if key not in seen:
                order.append(key)
            seen[key][value] += 1

    for key in order:
        values = seen[key]
        if len(values) > 1:
            result.sizes_disagree.append(key)
            continue
        (value,) = values
        result.edits.append(CellEdit(key[0], key[1], key[2], value))
    return result


def _table_from_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [row for _, row in zip(range(MAX_ROWS + _HEADER_SEARCH_ROWS), reader, strict=False)]


def parse_values_upload(
    filename: str, data: bytes, registry: dict[str, tuple[str, bool]]
) -> ParsedValues:
    """Dispatch on the extension. A workbook's sheets are tried in turn and
    the one yielding the most edits wins; a workbook with a `Values` or
    `SKU Detail` sheet is read from that sheet first."""
    lowered = (filename or "").lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            candidates: list[ParsedValues] = []
            sheets = sorted(
                workbook.worksheets,
                key=lambda ws: 0 if ws.title in ("Values", "SKU Detail") else 1,
            )
            for worksheet in sheets:
                table = [
                    [_norm(c) for c in row]
                    for _, row in zip(
                        range(MAX_ROWS + _HEADER_SEARCH_ROWS),
                        worksheet.iter_rows(values_only=True),
                        strict=False,
                    )
                ]
                if not table:
                    continue
                candidates.append(parse_table(table, registry, sheet=worksheet.title))
            usable = [c for c in candidates if c.edits or c.sizes_disagree]
            if usable:
                return max(usable, key=lambda c: len(c.edits))
            if candidates:
                return candidates[0]
            return ParsedValues(problem="The workbook has no readable sheet.")
        finally:
            workbook.close()
    if lowered.endswith((".csv", ".tsv", ".txt")):
        return parse_table(_table_from_csv(data), registry, sheet="the CSV")
    return ParsedValues(
        problem=f"{filename or 'The file'} is not a .xlsx or .csv. Save the sheet as one of those."
    )
