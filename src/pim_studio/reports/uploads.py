"""Turning somebody's file of ModelColors into a scope.

The input is whatever was to hand: a column pasted out of the order file, a
sheet the Turkey team sent back, a Check_list export, or forty codes typed into
a box. All of them are lists of product codes and none of them agree on how a
product code is written.

Two decisions shape this module.

**The catalogue is the oracle.** Nothing here parses a code into style and
colour. It finds the column most likely to hold ModelColors by *how many of its
values match products we hold* — so a sheet with `sku`, `Model`, `Colour` and a
`Barcode` column resolves without being told which one to read, and a sheet
whose codes are all unknown resolves to nothing rather than to a confident
reading of the wrong column.

**Unmatched is an output, not an error.** An upload of 500 codes that matches
380 has told you something about the other 120, and dropping them silently is
how a coverage number quietly stops covering. Every token that resolves to
nothing comes back in `not_found`, in the spelling it was written in.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from ..rules.identifiers import resolve_model_color, squash

#: Splits pasted text. Newlines, commas, semicolons and tabs are all used as
#: separators by somebody; spaces are not, because `L3558 132` is one code.
_PASTE_SPLIT = re.compile(r"[\n\r,;\t]+")

#: A cell that cannot be a product code however it is spelled. Cheap enough to
#: run over every cell of a 40-column sheet, and it keeps the column scoring
#: below from being swayed by a column of dates or prices.
_IMPOSSIBLE = re.compile(r"^[\d\s.,:/-]*$")

MAX_ROWS = 200_000


@dataclass
class ParsedUpload:
    """What a file turned out to contain.

    `column` and `sheet` are carried back so the panel can say *where* it read
    the codes from. A resolution the user cannot check is one they have to
    trust, and this is the step most likely to have picked the wrong column.
    """

    model_colors: list[str] = field(default_factory=list)
    not_found: list[str] = field(default_factory=list)
    #: Which column the codes came from, when the input had columns.
    column: str = ""
    sheet: str = ""
    rows_read: int = 0
    #: Codes given more than once and collapsed to one. Reported rather than
    #: absorbed: a list with duplicates in it is usually two sources pasted
    #: together, and the reader who sees 480 come back from 500 rows deserves
    #: to know which of the twenty were repeats and which were unreadable.
    duplicates: int = 0
    #: Set when nothing at all resolved, explaining what was looked at.
    problem: str = ""

    @property
    def matched(self) -> int:
        return len(self.model_colors)


def _index(known: list[str]) -> dict[str, str]:
    """`squash(code)` → canonical code, for the resolver to match against."""
    return {squash(code): code for code in known}


def _resolve_all(
    tokens: list[str], index: dict[str, str]
) -> tuple[list[str], list[str], int]:
    """Resolve in input order, de-duplicating both sides.

    Order is kept because an uploaded list is often already sorted the way the
    person wants to read the report, and re-sorting it makes them find their
    place again.
    """
    found: dict[str, None] = {}
    missing: dict[str, None] = {}
    seen = 0
    for token in tokens:
        text = str(token).strip()
        if not text:
            continue
        seen += 1
        match = resolve_model_color(text, index)
        if match:
            found.setdefault(match)
        else:
            missing.setdefault(text)
    # Counted against non-blank tokens, not `rows_read`: a sheet's blank cells
    # are not repeats, and deriving this from the row count would report them
    # as though they were.
    return list(found), list(missing), seen - len(found) - len(missing)


def parse_pasted(text: str, known: list[str]) -> ParsedUpload:
    """Codes typed or pasted into a box, one per line or comma-separated."""
    index = _index(known)
    tokens = [t.strip() for t in _PASTE_SPLIT.split(text or "") if t.strip()]
    matched, missing, repeats = _resolve_all(tokens, index)
    result = ParsedUpload(
        model_colors=matched,
        not_found=missing,
        rows_read=len(tokens),
        duplicates=repeats,
        column="pasted",
    )
    if not matched and tokens:
        result.problem = (
            f"None of the {len(tokens)} codes matched a product we hold. Codes "
            "match with or without dots, underscores or a size suffix, so this "
            "usually means the list is for a catalogue we have not synced."
        )
    return result


def _best_column(
    header: list[str], rows: list[list[str]], index: dict[str, str]
) -> tuple[int, int]:
    """The column index whose values resolve most often, and how many resolved.

    Scored rather than named. A header called `ModelColor` is good evidence and
    a header called `Unnamed: 3` is none, but a column where 480 of 500 values
    are products we hold is proof — and it is the only signal that survives a
    file whose headers are in Turkish, or missing, or lying.
    """
    best_index, best_score = -1, 0
    width = max((len(r) for r in rows), default=0)
    for column in range(max(width, len(header))):
        score = 0
        for row in rows:
            if column >= len(row):
                continue
            cell = str(row[column]).strip()
            if not cell or _IMPOSSIBLE.match(cell):
                continue
            if resolve_model_color(cell, index):
                score += 1
        if score > best_score:
            best_index, best_score = column, score
    return best_index, best_score


def _from_table(
    header: list[str], rows: list[list[str]], known: list[str], sheet: str = ""
) -> ParsedUpload:
    index = _index(known)
    column, score = _best_column(header, rows, index)
    if column < 0 or score == 0:
        return ParsedUpload(
            rows_read=len(rows),
            sheet=sheet,
            problem=(
                f"No column in {sheet or 'the file'} holds codes matching products "
                f"we hold — {len(rows)} rows and "
                f"{max((len(r) for r in rows), default=0)} columns were checked. "
                "Check the file identifies products by ModelColor or by a "
                "size-suffixed article code."
            ),
        )

    name = header[column] if column < len(header) else f"column {column + 1}"
    tokens = [row[column] for row in rows if column < len(row)]
    matched, missing, repeats = _resolve_all(tokens, index)
    return ParsedUpload(
        model_colors=matched,
        not_found=missing,
        column=name or f"column {column + 1}",
        sheet=sheet,
        rows_read=len(rows),
        duplicates=repeats,
    )


def parse_csv(data: bytes, known: list[str]) -> ParsedUpload:
    text = data.decode("utf-8-sig", errors="replace")
    # Sniffing rather than assuming a comma: the exports these lists are cut
    # from are semicolon-delimited when they come out of a Polish Excel.
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    table = [row for _, row in zip(range(MAX_ROWS + 1), reader, strict=False)]
    if not table:
        return ParsedUpload(problem="The file is empty.")
    return _from_table(table[0], table[1:], known, sheet="the CSV")


#: Returned when a workbook has no sheet with any rows at all — as opposed to
#: one whose rows simply hold nothing we recognise, which is a different
#: message and a far more common one.
_NO_SHEET = "The workbook has no readable sheet."


def _best_of(candidates: list[ParsedUpload]) -> ParsedUpload:
    """The sheet that resolved the most codes, or the fullest one it read.

    The fallback matters more than the winner. Picking by `matched` alone left
    a workbook where nothing resolved reporting "no readable sheet" — which
    says we could not open the file, when in fact we read every row of it and
    recognised none of the codes. Those send the user to two different places.
    """
    if not candidates:
        return ParsedUpload(problem=_NO_SHEET)
    resolved = [c for c in candidates if c.matched]
    if resolved:
        return max(resolved, key=lambda c: c.matched)
    return max(candidates, key=lambda c: c.rows_read)


def parse_workbook(data: bytes, known: list[str]) -> ParsedUpload:
    """An .xlsx. Every sheet is tried; the one that resolves most codes wins.

    A list of ModelColors is rarely on the first sheet of the workbook it
    arrives in — it is on `Sheet2`, or on the one named after the season, with
    a cover sheet in front of it.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        candidates = []
        for worksheet in workbook.worksheets:
            table = [
                ["" if c is None else str(c) for c in row]
                for _, row in zip(
                    range(MAX_ROWS + 1),
                    worksheet.iter_rows(values_only=True),
                    strict=False,
                )
            ]
            if not table:
                continue
            candidates.append(
                _from_table(table[0], table[1:], known, sheet=worksheet.title)
            )
        return _best_of(candidates)
    finally:
        workbook.close()


def parse_legacy_xls(data: bytes, known: list[str]) -> ParsedUpload:
    """The .xls the backoffice still exports. `xlrd` is a hard dependency here
    for exactly this reason — the Check_list export has never been .xlsx."""
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    candidates = []
    for sheet in book.sheets():
        table = [
            [str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            for r in range(min(sheet.nrows, MAX_ROWS + 1))
        ]
        if not table:
            continue
        candidates.append(_from_table(table[0], table[1:], known, sheet=sheet.name))
    return _best_of(candidates)


def parse_upload(filename: str, data: bytes, known: list[str]) -> ParsedUpload:
    """Dispatch on the extension, since the panel accepts all three."""
    lowered = (filename or "").lower()
    if lowered.endswith(".xlsx") or lowered.endswith(".xlsm"):
        return parse_workbook(data, known)
    if lowered.endswith(".xls"):
        return parse_legacy_xls(data, known)
    if lowered.endswith(".csv") or lowered.endswith(".tsv") or lowered.endswith(".txt"):
        return parse_csv(data, known)
    raise ValueError(
        f"Cannot read {filename!r}. Upload a .xlsx, .xls, .csv or .txt, or paste "
        "the codes instead."
    )
