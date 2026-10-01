"""The handoff — a ZIP of the photos and an Excel that says where they go.

This is the module's actual output. Our Omnitron connection is read-only by
decision, so no amount of publishing puts a photo on a product page: somebody
on the Turkey team has to upload it. A run that ends with files on the photo
server and nothing for them to act on has not finished, it has stopped.

**The Excel is the operative half.** Omnitron imports images by URL, so the
sheet — one row per photo, with its ModelColor and its public address — is what
actually does the work. It is written in the same style as the coverage
workbook, because it lands in the same inbox and a file that looks like a
different tool's output gets treated as one.

**The ZIP is the fallback**, for the case where fetching from our host is
refused or slow. It costs disk and nothing else, and the one time it is needed
it saves a week.

**A bundle's row outlives its files.** Bundles are pruned after
`settings.bundles_kept`; "what did we hand over in August" has to stay
answerable long after the ZIP is gone, so pruning clears the files and marks
the row `cleared` rather than deleting it.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import integrations
from ..logging import log
from ..models import HandoffBundle, PublishedImage, utcnow

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
STRIPE_FILL = PatternFill("solid", fgColor="F7F9FC")
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="FFFFFF")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
NOTE_FONT = Font(name=FONT, size=9, color="595959")

COLUMNS = [
    ("ModelColor", 18),
    ("Edition", 12),
    ("File name", 34),
    ("Image URL", 74),
    ("Published", 20),
]

#: How the sheet explains itself. The file is emailed and read by people who
#: were not in the conversation that produced it, so it says what it is for.
BLURB = (
    "Product photos for options that have no image in Omnitron. "
    "Each row is one photo and its public address — Omnitron can import "
    "images by URL, so the ZIP beside this file is only needed if fetching "
    "from that address is refused. Original and Zalando editions are listed "
    "separately; an option may have several photos."
)


@dataclass(frozen=True, slots=True)
class BundleFiles:
    bundle_id: int
    xlsx: Path
    zip: Path | None
    options: int
    files: int
    bytes: int


def pending(db: Session) -> list[PublishedImage]:
    """Published photos not yet in any bundle.

    `bundle_id IS NULL` is the queue. A photo joins exactly one bundle, so a
    second run does not re-send what the last one already handed over — and
    re-sending is not harmless: it reads to the recipient as a correction.
    """
    return list(
        db.scalars(
            select(PublishedImage)
            .where(PublishedImage.bundle_id.is_(None))
            .order_by(PublishedImage.model_color, PublishedImage.asset_name)
        )
    )


def build(
    db: Session,
    root: Path,
    built_by: str | None = None,
    with_zip: bool = True,
    fetch: bool = True,
) -> BundleFiles | None:
    """Build a bundle from everything published and not yet handed over.

    Returns None when there is nothing to send — which is a normal weekly
    outcome and must not produce an empty ZIP somebody has to open to discover
    it is empty.

    `fetch` downloads each photo back from the photo server to fill the ZIP.
    Off in tests, and off when only the Excel is wanted.
    """
    rows = pending(db)
    if not rows:
        log.info("bundle_nothing_to_send")
        return None

    bundle = HandoffBundle(
        built_by=built_by,
        option_count=len({r.model_color for r in rows}),
        file_count=len(rows),
        bytes=sum(r.bytes or 0 for r in rows),
        status="building",
    )
    db.add(bundle)
    db.commit()

    stamp = bundle.built_at.strftime("%Y-%m-%d") if bundle.built_at else "bundle"
    directory = root / f"{stamp}-{bundle.id:04d}"
    directory.mkdir(parents=True, exist_ok=True)

    xlsx = directory / f"northwind-images-{stamp}.xlsx"
    _workbook(rows, bundle).save(xlsx)

    archive: Path | None = None
    if with_zip:
        archive = directory / f"northwind-images-{stamp}.zip"
        _archive(rows, archive, fetch=fetch)

    # Claim the rows only once the files exist. A bundle that failed half way
    # through must not leave its photos marked as handed over — they would
    # never appear in another bundle and would silently never be sent.
    for row in rows:
        row.bundle_id = bundle.id
    bundle.status = "ready"
    bundle.zip_path = str(archive) if archive else None
    bundle.xlsx_path = str(xlsx)
    db.commit()

    log.info(
        "bundle_built",
        bundle=bundle.id,
        options=bundle.option_count,
        files=bundle.file_count,
    )
    return BundleFiles(
        bundle_id=bundle.id,
        xlsx=xlsx,
        zip=archive,
        options=bundle.option_count,
        files=bundle.file_count,
        bytes=bundle.bytes,
    )


def prune(db: Session, root: Path, keep: int) -> int:
    """Delete the files of all but the newest `keep` bundles. Returns how many.

    The rows stay and become `cleared`. What was handed over in August is a
    question that outlives the ZIP by years.
    """
    del root  # paths are on the rows; the root is only for building
    stale = list(
        db.scalars(
            select(HandoffBundle)
            .where(HandoffBundle.status == "ready")
            .order_by(HandoffBundle.built_at.desc())
            .offset(keep)
        )
    )
    for bundle in stale:
        for path in (bundle.zip_path, bundle.xlsx_path):
            if path:
                Path(path).unlink(missing_ok=True)
        directory = Path(bundle.xlsx_path).parent if bundle.xlsx_path else None
        if directory and directory.exists() and not any(directory.iterdir()):
            directory.rmdir()
        bundle.status = "cleared"
    if stale:
        db.commit()
        log.info("bundles_pruned", count=len(stale), kept=keep)
    return len(stale)


def mark_sent(db: Session, bundle_id: int, when: datetime | None = None) -> None:
    """Record that somebody sent this on. Null means built and not yet sent."""
    bundle = db.get(HandoffBundle, bundle_id)
    if bundle is None:
        return
    bundle.sent_at = when or utcnow()
    db.commit()


def _workbook(rows: list[PublishedImage], bundle: HandoffBundle) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Images"

    width = len(COLUMNS)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    title = sheet.cell(row=1, column=1, value="Northwind CEE — product images")
    title.fill = TITLE_FILL
    title.font = TITLE_FONT
    title.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 26

    sheet.merge_cells(start_row=2, start_column=1, end_row=3, end_column=width)
    note = sheet.cell(row=2, column=1, value=BLURB)
    note.font = NOTE_FONT
    note.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
    sheet.row_dimensions[2].height = 28

    stamp = bundle.built_at.strftime("%d %B %Y") if bundle.built_at else ""
    meta = sheet.cell(
        row=4,
        column=1,
        value=(
            f"{bundle.option_count} options · {bundle.file_count} photos · built {stamp}"
        ),
    )
    meta.font = NOTE_FONT

    header_row = 6
    for index, (label, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", indent=1)

    for offset, row in enumerate(rows):
        line = header_row + 1 + offset
        published = row.published_at.strftime("%Y-%m-%d %H:%M") if row.published_at else ""
        values = [
            row.model_color,
            row.edition,
            row.remote_name,
            row.public_url,
            published,
        ]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=line, column=index, value=value)
            cell.font = BODY_FONT
            if offset % 2:
                cell.fill = STRIPE_FILL

    for index, (_, size) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = size

    # Freeze under the header so the columns stay named while somebody scrolls
    # a thousand rows — this file is read, not glanced at.
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(width)}{header_row + len(rows)}"
    )
    return workbook


def _archive(rows: list[PublishedImage], destination: Path, fetch: bool = True) -> None:
    """Fetch each photo back from the photo server into a ZIP.

    Read back rather than kept from the upload: the run that published them may
    have been weeks ago, and staging a gigabyte of photos against a bundle
    nobody has asked for yet is the wrong trade. A file that cannot be fetched
    is skipped with a warning — a bundle missing one photo is worth sending,
    and a bundle that refuses to build over one is not.
    """
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for row in rows:
            if not fetch:
                archive.writestr(row.remote_name, b"")
                continue
            try:
                response = integrations.http_get(
                    row.public_url, timeout=60.0, follow_redirects=True
                )
                response.raise_for_status()
            except httpx.HTTPError as exc:
                log.warning(
                    "bundle_photo_unfetchable",
                    remote=row.remote_name,
                    error=str(exc)[:160],
                )
                continue
            archive.writestr(row.remote_name, response.content)
